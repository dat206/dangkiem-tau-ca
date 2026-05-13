import re
import docx
from pydantic import BaseModel, ValidationError

class ParseError(Exception):
    pass

class VesselData(BaseModel):
    so_dang_ky: str
    ma_tinh: str
    lmax: float
    hinh_thuc_kiem_tra: str
    cap_tau: str

INSPECTION_TYPES = {
    "ĐK": "Định kỳ",
    "HN": "Hàng năm",
    "TĐ": "Trên đà",
    "GS": "Giám sát",
    "CH": "Cải hoán"
}

def clean_text(text: str) -> str:
    text = re.sub(r'\s+', ' ', text)
    text = text.replace('\n', ' ').strip()
    return text

def parse_vessel_docx(file_path: str) -> VesselData:
    try:
        doc = docx.Document(file_path)
    except Exception as e:
        raise ParseError(f"Không thể đọc file DOCX: {str(e)}")

    full_text_blocks = []
    tables_data = []

    for para in doc.paragraphs:
        full_text_blocks.append(clean_text(para.text))

    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_data = [clean_text(cell.text) for cell in row.cells]
            table_rows.append(row_data)
            full_text_blocks.extend(row_data)
        tables_data.append(table_rows)

    full_text = " | ".join(full_text_blocks)

    so_dang_ky_match = re.search(r'Số đăng ký:\s*([A-ZĐ]+-\d+-TS)', full_text, re.IGNORECASE)
    if not so_dang_ky_match:
        raise ParseError("Không tìm thấy Số đăng ký")
    so_dang_ky = so_dang_ky_match.group(1).strip()

    ma_tinh_match = re.match(r'^([A-ZĐ]+)-', so_dang_ky)
    ma_tinh = ma_tinh_match.group(1) if ma_tinh_match else ""

    lmax_match = re.search(r'Chiều dài,\s*Lmax:\s*([\d,]+)', full_text, re.IGNORECASE)
    if not lmax_match:
        raise ParseError("Không tìm thấy Chiều dài Lmax")
    lmax_str = lmax_match.group(1).replace(',', '.')
    try:
        lmax = float(lmax_str)
    except ValueError:
        raise ParseError(f"Định dạng Lmax không hợp lệ: {lmax_str}")

    hinh_thuc_kiem_tra = "Không xác định"
    so_chung_nhan_match = re.search(r'Số:\s*[\d\.]+/(ĐK|HN|TĐ|GS|CH)', full_text)
    if so_chung_nhan_match:
        ma_hk = so_chung_nhan_match.group(1)
        hinh_thuc_kiem_tra = INSPECTION_TYPES.get(ma_hk, ma_hk)

    cap_tau = "Không xác định"
    for table in tables_data:
        if len(table) >= 2 and "Cấp tàu" in table[0][0]:
            header_row = table[0]
            value_row = table[1]
            
            cap_tau_cols = [
                "Hạn chế III",
                "Hạn chế II",
                "Hạn chế I",
                "Không hạn chế"
            ]
            
            for cap_name in cap_tau_cols:
                for idx, cell_val in enumerate(header_row):
                    if cap_name in cell_val:
                        if idx < len(value_row) and value_row[idx].strip().upper() == "X":
                            cap_tau = cap_name
                            break
                if cap_tau != "Không xác định":
                    break

    try:
        return VesselData(
            so_dang_ky=so_dang_ky,
            ma_tinh=ma_tinh,
            lmax=lmax,
            hinh_thuc_kiem_tra=hinh_thuc_kiem_tra,
            cap_tau=cap_tau
        )
    except ValidationError as e:
        raise ParseError(f"Lỗi xác thực dữ liệu: {str(e)}")